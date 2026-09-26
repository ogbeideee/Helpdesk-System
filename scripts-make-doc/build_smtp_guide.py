import datetime
import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUTPUT = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "M365-SMTP-OAuth-Setup-and-Verification-Guide.docx")

NAVY = RGBColor(0x1F, 0x3B, 0x66)
BLUE = RGBColor(0x2E, 0x74, 0xB5)
GREEN = RGBColor(0x1E, 0x6B, 0x3A)
RED = RGBColor(0x9C, 0x27, 0x27)
GREY = RGBColor(0x55, 0x55, 0x55)
CODE_FILL = "F2F2F2"
NOTE_FILL = "EAF3FB"
WARN_FILL = "FDECEA"
OK_FILL = "E8F5E9"


def shade(cell_or_paragraph, fill):
    if hasattr(cell_or_paragraph, "_tc"):
        properties = cell_or_paragraph._tc.get_or_add_tcPr()
    else:
        properties = cell_or_paragraph._p.get_or_add_pPr()
    element = OxmlElement("w:shd")
    element.set(qn("w:val"), "clear")
    element.set(qn("w:fill"), fill)
    properties.append(element)


def set_cell_text(cell, text, bold=False, color=None, size=9.5):
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(str(text))
    run.bold = bold
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    return paragraph


def add_table(document, headers, rows, widths=None, header_color=BLUE):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], header, bold=True, color=header_color)
        shade(table.rows[0].cells[index], "D9EAF7")
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            set_cell_text(cells[index], value)
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Inches(width)
    document.add_paragraph()
    return table


def add_code(document, value):
    lines = str(value).split("\n")
    for index, line in enumerate(lines):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.25)
        paragraph.paragraph_format.space_after = Pt(0 if index < len(lines) - 1 else 6)
        run = paragraph.add_run(line if line else " ")
        run.font.name = "Consolas"
        run.font.size = Pt(9)
        shade(paragraph, CODE_FILL)
    return paragraph


def add_paragraph(document, text="", bold=False, italic=False, color=None, size=None):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.name = "Calibri"
    if color:
        run.font.color.rgb = color
    if size:
        run.font.size = Pt(size)
    return paragraph


def add_bullet(document, text, level=0):
    paragraph = document.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.add_run(text)
    return paragraph


def add_number(document, text):
    paragraph = document.add_paragraph(style="List Number")
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.add_run(text)
    return paragraph


def add_callout(document, label, text, fill=NOTE_FILL, color=NAVY):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.1)
    paragraph.paragraph_format.right_indent = Inches(0.1)
    label_run = paragraph.add_run(label + "  ")
    label_run.bold = True
    label_run.font.color.rgb = color
    paragraph.add_run(text)
    shade(paragraph, fill)
    return paragraph


def add_page_break(document):
    document.add_page_break()


def add_heading(document, text, level):
    heading = document.add_heading(text, level=level)
    heading.paragraph_format.keep_with_next = True
    return heading


def add_checklist(document, items):
    for item in items:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.1)
        paragraph.paragraph_format.space_after = Pt(3)
        checkbox = paragraph.add_run("☐  ")
        checkbox.font.size = Pt(12)
        if isinstance(item, tuple):
            label = paragraph.add_run(item[0])
            label.bold = True
            paragraph.add_run(item[1])
        else:
            paragraph.add_run(item)


def configure_document(document):
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    settings = [(20, NAVY, 8), (15, NAVY, 14), (12.5, BLUE, 10), (11.5, BLUE, 8)]
    for level, (size, color, before) in enumerate(settings, start=1):
        style = document.styles["Heading %d" % level]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.keep_with_next = True

    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("Internal IT Operations  •  M365 SMTP OAuth Setup Guide")
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = GREY


def add_header_row_emphasis(document, text):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(text)
    run.bold = True
    run.font.color.rgb = NAVY
    return paragraph


def build_document():
    document = Document()
    configure_document(document)
    properties = document.core_properties
    properties.title = "M365 SMTP OAuth Setup and Verification Guide"
    properties.subject = "Sending TicketDesk notifications from a Microsoft 365 shared mailbox without Microsoft Graph"
    properties.author = "TicketDesk"
    properties.keywords = "Microsoft 365, SMTP, XOAUTH2, shared mailbox, helpdesk, OAuth"
    properties.comments = "Implementation and verification runbook."

    for _ in range(5):
        document.add_paragraph()
    title = add_paragraph(document, "M365 SMTP OAuth Setup and Verification Guide", bold=True, size=28, color=NAVY)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = add_paragraph(document, "Sending TicketDesk notifications from the Microsoft 365 helpdesk shared mailbox", size=14, color=BLUE)
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph()
    summary = add_paragraph(document, "Implementation, Microsoft 365 configuration, application wiring, security controls, testing, production rollout, and troubleshooting", italic=True, size=11)
    summary.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for _ in range(3):
        document.add_paragraph()
    meta = add_paragraph(document, "Prepared: %s    |    Version 1.0    |    Classification: Internal – IT Operations" % datetime.date.today().strftime("%d %B %Y"), size=10)
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_callout(document, "IMPORTANT SCOPE:", "This guide does not claim that the repository already sends through SMTP. The current application sends through Microsoft Graph when configured and otherwise logs notifications to the console. The steps below are the implementation and Microsoft 365 setup runbook for adding an SMTP OAuth transport.", WARN_FILL, RED)
    add_page_break(document)

    add_heading(document, "1. Purpose and decision", 1)
    add_paragraph(document, "The required mail path is:")
    add_code(document, "Requester email → Microsoft 365 helpdesk shared mailbox → forwarding rule → Gmail mailbox → IMAP poller → TicketDesk ticket\nTicket resolved or closed → TicketDesk → M365 SMTP AUTH/OAuth → requester")
    add_paragraph(document, "The forwarding rule is for inbound mail. It does not make Gmail the sender of TicketDesk notifications. The outbound connection must authenticate to Microsoft 365 and submit the message as the helpdesk shared mailbox.")
    add_callout(document, "NO GRAPH API REQUIRED:", "Direct SMTP does not call the Microsoft Graph mail API. It does, however, use Microsoft Entra ID for the application identity and Exchange Online for the shared-mailbox SendAs permission. Those are identity and mailbox permissions, not Graph API traffic.", OK_FILL, GREEN)
    add_paragraph(document, "This guide uses unattended client-credentials authentication. It does not use a shared-mailbox password, an app password, a delegated user sign-in, or Gmail as the outbound sender.")

    add_heading(document, "2. Current repository capability", 1)
    add_paragraph(document, "The existing mail architecture is suitable for adding this transport without creating a second notification system:")
    add_table(document, ["Current behavior", "Evidence", "Required change"], [
        ["Ticket resolution and closure call the requester status notification", "server/routes/tickets.js:451-453", "No lifecycle change; the selected transport receives the call."],
        ["Status and resolution wording is already built", "server/src/email/outbound.js:184-213", "No wording change is required for the first SMTP implementation."],
        ["All outbound notification functions go through the shared mailer", "server/src/mailer.js:82-167", "Extend transport selection and preserve safe-send behavior."],
        ["The default transport is Graph or console only", "server/src/mailer.js:67-72", "Add SMTP transport selection before Graph and console fallback."],
        ["No SMTP package is installed", "server/package.json:65-75", "Add and approve a tested SMTP client dependency."],
        ["The IMAP poller already handles forwarded Gmail mail", "server/src/imap/mailService.js:205-274", "Keep this inbound path; no Gmail sending configuration is required."],
    ], [2.0, 2.2, 2.6])
    add_callout(document, "IMPLEMENTATION CONSEQUENCE:", "Adding SMTP is primarily a transport/configuration change. The current application still needs a new transport module, OAuth token acquisition, environment documentation, and tests. Merely adding SMTP variables to .env will not make mail work.", WARN_FILL, RED)

    add_heading(document, "3. Roles and prerequisites", 1)
    add_paragraph(document, "Complete these items before configuration. The person operating the helpdesk may not have all required permissions, so identify the Microsoft 365 administrator first.")
    add_checklist(document, [
        ("Microsoft 365 administrator access: ", "ability to register an Entra application, grant application permission consent, connect to Exchange Online PowerShell, and manage mailbox permissions."),
        ("Exchange Online administrator access: ", "ability to inspect Authenticated SMTP settings and grant the application service principal SendAs permission."),
        ("A working shared mailbox: ", "for example, helpdesk@company.com, with a valid primary SMTP address and an established inbound forwarding rule to the Gmail mailbox used by IMAP."),
        ("A verified application deployment target: ", "the server environment where TicketDesk runs, such as Fly.io, plus a secure way to set environment secrets."),
        ("A test recipient: ", "an address outside the helpdesk mailbox and outside the forwarding path, so delivery and threading can be tested safely."),
        ("Network access: ", "outbound TCP 587 from the TicketDesk host to smtp.office365.com, with TLS support."),
    ])

    add_heading(document, "4. Target architecture and security boundary", 1)
    add_paragraph(document, "The two email directions are deliberately separate:")
    add_table(document, ["Direction", "Server", "Authentication", "Data source or recipient"], [
        ["Inbound", "Microsoft 365 shared mailbox → Gmail forwarding → IMAP", "Existing Gmail IMAP password or XOAUTH2 configuration", "Gmail mailbox"],
        ["Outbound", "TicketDesk → smtp.office365.com:587", "OAuth 2.0 client credentials plus SMTP XOAUTH2", "Requester email address"],
    ], [1.1, 3.0, 2.0, 1.4])
    add_paragraph(document, "The SMTP connection uses STARTTLS. The application secret is used only to obtain a short-lived OAuth access token from Microsoft Entra. The access token is placed in the SMTP XOAUTH2 exchange and is never written to the database, sent to the browser, or printed in logs.")
    add_callout(document, "SENDER RULE:", "The SMTP envelope sender and visible From address must both be the Microsoft 365 shared mailbox. Set Reply-To to the same shared mailbox so a requester's reply returns through the O365 forwarding path. Do not set From or Reply-To to the Gmail address.", NOTE_FILL)

    add_heading(document, "5. Microsoft 365 configuration", 1)
    add_paragraph(document, "Use the Microsoft Entra admin center and Exchange Online PowerShell. The exact portal labels can change, but the objects and permissions must exist.")

    add_heading(document, "5.1 Confirm the shared mailbox and forwarding", 2)
    add_number(document, "Confirm the mailbox exists, is not disabled, and has the expected primary SMTP address. The application must send from the exact address that Exchange recognizes as the shared mailbox.")
    add_number(document, "Confirm the existing forwarding rule delivers incoming requester mail to the Gmail account used by the IMAP poller. Do not change that forwarding rule for SMTP.")
    add_number(document, "Confirm the Gmail mailbox is still configured only for inbound IMAP. Gmail must not receive the SMTP sender credentials and must not be used as the outbound From address.")
    add_number(document, "Choose a test recipient outside the forwarding path. A requester reply to the shared mailbox should be forwarded to Gmail and become ticket activity, but the test message used to verify outbound delivery should go to the external test address.")

    add_heading(document, "5.2 Create a dedicated Entra application", 2)
    add_paragraph(document, "Create a dedicated application registration, rather than reusing the existing Graph application or a human user account. The dedicated registration makes permission review, secret rotation, and future isolated deployments possible.")
    add_number(document, "In the Microsoft Entra admin center, open App registrations and create an application named, for example, TicketDesk SMTP Sender.")
    add_number(document, "Use the single-tenant option for the organisation that owns the helpdesk mailbox.")
    add_number(document, "Create a client secret for the first deployment. For production, a certificate credential is preferable because it avoids a long-lived password-like secret; whichever credential is selected must be stored only in the deployment secret store.")
    add_number(document, "Record the Tenant ID, Application (client) ID, and the secret or certificate details in the deployment secret manager. Do not put them in the repository, client bundle, screenshots, tickets, or chat messages.")

    add_heading(document, "5.3 Grant SMTP application permission and admin consent", 2)
    add_paragraph(document, "In the application registration, add the Office 365 Exchange Online application permission SMTP.SendAsApp. The Microsoft documentation for client-credentials SMTP uses this permission and the scope https://outlook.office.com/.default. It is not the Graph Mail.ReadWrite or Mail.Send permission.")
    add_table(document, ["Setting", "Required value", "Why"], [
        ["Permission type", "Application permissions", "The server runs unattended; it cannot depend on an administrator signing in."],
        ["Exchange Online permission", "SMTP.SendAsApp", "Allows the service principal to submit mail through Exchange SMTP."],
        ["Admin consent", "Granted by a Microsoft 365 administrator", "Application permissions require tenant administrator consent."],
        ["OAuth token scope", "https://outlook.office.com/.default", "Requests the Exchange Online SMTP application permission from the tenant."],
    ], [1.5, 2.4, 3.0])
    add_callout(document, "DO NOT CONFUSE PERMISSIONS:", "SMTP.SendAsApp is different from the Graph permission Graph Mail.Send. The former authorizes SMTP submission; the latter authorizes Graph API calls. The requested setup uses the former and does not call Graph.", NOTE_FILL)

    add_heading(document, "5.4 Register the application service principal in Exchange Online", 2)
    add_paragraph(document, "A client-credentials SMTP application must be registered as an Exchange service principal. Connect to Exchange Online PowerShell using an administrator account, then run the equivalent of the following commands. Replace placeholders; do not paste real values into a shared document.")
    add_code(document, "Connect-ExchangeOnline -Interactive\n\n$ClientId = '<APPLICATION_CLIENT_ID>'\n$SharedMailbox = 'helpdesk@company.com'\n\n$EntraSp = Get-AzureADServicePrincipal -Filter \"appId eq '$ClientId'\"\n$EntraSp | Format-List AppId, ObjectId, DisplayName\n\n$ExistingExchangeSp = Get-ServicePrincipal -Identity $ClientId -ErrorAction SilentlyContinue\nif (-not $ExistingExchangeSp) {\n    New-ServicePrincipal -AppId $ClientId -ObjectId $EntraSp.ObjectId\n}\n\n$ExchangeSp = Get-ServicePrincipal -Identity $ClientId\n$ExchangeSp | Format-List Identity, ObjectId, DisplayName")
    add_paragraph(document, "If Get-AzureADServicePrincipal is unavailable in the administrator's environment, use the equivalent Microsoft Entra service-principal lookup available to that environment. The required values are the application AppId and the Entra service-principal ObjectId. The ObjectId from the App registration overview is not necessarily the Exchange service-principal identity; verify the values returned by Get-ServicePrincipal.")
    add_callout(document, "IMPORTANT:", "If the New-ServicePrincipal command reports that the service principal already exists, do not create duplicates. Continue by retrieving the existing Exchange service principal and verify that it belongs to the intended Entra application.", NOTE_FILL)

    add_heading(document, "5.5 Grant the service principal SendAs on the shared mailbox", 2)
    add_paragraph(document, "Grant only the permission needed for this use case. TicketDesk needs to send as the shared mailbox, not necessarily to read the mailbox through this service principal.")
    add_code(document, "Add-RecipientPermission \\\n  -Identity $SharedMailbox \\\n  -User $ExchangeSp.Identity \\\n  -AccessRights SendAs \\\n  -Confirm:$false\n\nGet-RecipientPermission -Identity $SharedMailbox | Format-Table -AutoSize")
    add_paragraph(document, "The resulting permission list must show the Exchange service principal with SendAs. If SendAs is missing, the SMTP server may authenticate the service principal but reject the message because it is not permitted to use the sender address.")
    add_callout(document, "NO SHARED-MAILBOX PASSWORD:", "Do not reset the shared mailbox password and do not use it as SMTP_PASSWORD. The shared mailbox is not a normal interactive login account, and the recommended design is OAuth client credentials with service-principal SendAs.", WARN_FILL, RED)

    add_heading(document, "5.6 Check Authenticated SMTP policy", 2)
    add_paragraph(document, "Direct SMTP submission uses the SMTP AUTH protocol even when the credential is OAuth rather than a password. Confirm that Authenticated SMTP is allowed for the shared mailbox and that the organisation's authentication policy does not block OAuth SMTP submission.")
    add_code(document, "Get-TransportConfig | Format-List SmtpClientAuthenticationDisabled\nGet-CASMailbox -Identity $SharedMailbox | Format-List PrimarySmtpAddress, SmtpClientAuthenticationDisabled")
    add_paragraph(document, "If the organisation uses Security Defaults or another authentication policy that disables SMTP AUTH, the Microsoft 365 administrator must resolve that policy deliberately. Do not broadly weaken tenant security controls merely to make the application work. Prefer enabling Authenticated SMTP only for the required mailbox where the tenant policy permits it, and document the approved exception.")
    add_callout(document, "2026 AUTHENTICATION NOTE:", "Microsoft has retired basic-password authentication for client SMTP submission. The implementation in this guide uses OAuth 2.0/XOAUTH2, not username/password SMTP authentication.", WARN_FILL, RED)

    add_heading(document, "6. TicketDesk application configuration", 1)
    add_paragraph(document, "The following configuration is the intended application contract. The current repository does not yet contain these variables or a matching transport; the implementation must add them before this section is operational.")
    add_table(document, ["Variable", "Example or format", "Secret", "Purpose"], [
        ["SMTP_ENABLED", "true", "No", "Enables the SMTP transport independently of Graph."],
        ["SMTP_HOST", "smtp.office365.com", "No", "Microsoft 365 SMTP submission host."],
        ["SMTP_PORT", "587", "No", "STARTTLS submission port."],
        ["SMTP_SECURE", "false", "No", "Use false with STARTTLS on port 587; do not use implicit TLS on 587."],
        ["SMTP_TENANT_ID", "<Entra tenant ID>", "No", "Authority used to obtain the access token."],
        ["SMTP_CLIENT_ID", "<Application client ID>", "No", "Application identity for client credentials."],
        ["SMTP_CLIENT_SECRET", "<client secret or certificate path>", "Yes", "Credential for obtaining the short-lived access token. Never log it."],
        ["SMTP_SHARED_MAILBOX", "helpdesk@company.com", "No", "Shared mailbox identity used for the XOAUTH2 username and SendAs checks."],
        ["SMTP_FROM", "helpdesk@company.com", "No", "Visible From address and envelope sender."],
        ["SMTP_REPLY_TO", "helpdesk@company.com", "No", "Reply address so replies return to O365 and are forwarded to Gmail."],
        ["SMTP_BROADCAST_DL", "Optional internal distribution list", "No", "Target for new-ticket team broadcasts; only if that notification is required."],
        ["SMTP_OAUTH_SCOPE", "https://outlook.office.com/.default", "No", "OAuth scope for Exchange SMTP client credentials."],
    ], [1.65, 2.15, 0.6, 2.7])
    add_callout(document, "NO SMTP_PASSWORD:", "The recommended configuration has no SMTP_PASSWORD. Do not add a shared-mailbox password or a Gmail password to this configuration.", WARN_FILL, RED)

    add_heading(document, "6.1 Recommended file changes", 2)
    add_table(document, ["File", "Required work"], [
        ["server/src/smtp/config.js", "Read and validate SMTP variables without returning secrets. Keep SMTP configuration separate from Graph configuration."],
        ["server/src/smtp/oauth.js", "Acquire and cache the Exchange SMTP access token using the existing @azure/msal-node dependency and the SMTP scope. Refresh before expiry and on a 401."],
        ["server/src/smtp/transport.js", "Implement sendMail(mail), sendBroadcastMail(mail), and hasBroadcastTarget(). Connect to smtp.office365.com:587, require STARTTLS, authenticate with SASL XOAUTH2, and send the existing mail shape."],
        ["server/src/mailer.js", "Select SMTP first when SMTP_ENABLED=true, Graph second when configured, and console last. Keep sendMailSafe and the existing notification entry points."],
        ["server/src/mailer.js", "Move broadcast-target policy away from GRAPH_BROADCAST_DL so an SMTP deployment does not depend on Graph variables."],
        ["server/.env.example", "Document SMTP settings with empty secrets and no organisation-specific mailbox address."],
        ["server/scripts/test-smtp.js", "Add unit and live verification tests using an injected fake transport and an explicit opt-in live probe."],
    ], [2.2, 4.7])
    add_paragraph(document, "The SMTP transport must convert the existing mailer shape without changing the email builders. The current shape contains subject, body, toRecipients, and optionally ccRecipients. The transport should map these fields to the SMTP client and set From, Reply-To, Date, Message-ID, and MIME headers. It must not add internal notes, audit data, or credentials to a requester message.")

    add_heading(document, "6.2 Transport selection and failure behaviour", 2)
    add_paragraph(document, "The selection order should be deterministic:")
    add_number(document, "If SMTP_ENABLED=true and the SMTP configuration is complete, use the SMTP transport.")
    add_number(document, "Otherwise, if the existing Graph configuration is enabled, use the Graph transport for backward compatibility.")
    add_number(document, "Otherwise, use the console transport for local development.")
    add_paragraph(document, "SMTP delivery must preserve the current fire-and-forget ticket behaviour: a ticket transition must not be rolled back because the external mail provider is temporarily unavailable. The mailer should log a sanitized error and return false. For reliable production reporting, add an outbox or Notification delivery record with retry status rather than relying only on process logs.")
    add_callout(document, "DO NOT LOG SECRETS:", "Log only the transport type, recipient, ticket number, outcome, and sanitized error. Never log the client secret, access token, XOAUTH2 payload, full MIME body, or full inbound email body.", WARN_FILL, RED)

    add_heading(document, "6.3 Set secrets on Fly.io", 2)
    add_paragraph(document, "For the current Fly.io deployment, set secrets with fly secrets rather than committing them to server/.env or any source file. The command below shows names and placeholders only:")
    add_code(document, "fly secrets set \\\n  SMTP_ENABLED=true \\\n  SMTP_HOST=smtp.office365.com \\\n  SMTP_PORT=587 \\\n  SMTP_SECURE=false \\\n  SMTP_TENANT_ID='<ENTRA_TENANT_ID>' \\\n  SMTP_CLIENT_ID='<SMTP_APP_CLIENT_ID>' \\\n  SMTP_CLIENT_SECRET='<SMTP_APP_CLIENT_SECRET>' \\\n  SMTP_SHARED_MAILBOX='helpdesk@company.com' \\\n  SMTP_FROM='helpdesk@company.com' \\\n  SMTP_REPLY_TO='helpdesk@company.com' \\\n  SMTP_BROADCAST_DL='' \\\n  SMTP_OAUTH_SCOPE='https://outlook.office.com/.default' \\\n  -a <APPLICATION_NAME>")
    add_paragraph(document, "The deployment must be restarted after secrets change. Confirm the variable names with fly secrets list, but remember that secret values are hidden. For multiple entities using the isolated-deployment model, repeat this configuration separately for each entity and never reuse one entity's client secret or database.")
    add_callout(document, "ROTATION:", "Create the new client secret, deploy and test it, then remove the old secret from Entra. Keep at least one known-good credential during rotation. A certificate credential can be used instead when the SMTP implementation supports it.", NOTE_FILL)

    add_heading(document, "7. End-to-end functional verification", 1)
    add_paragraph(document, "Perform these tests in order. Do not declare SMTP production-ready until the first five tests pass with real Microsoft 365 credentials and a real external test recipient.")
    add_checklist(document, [
        "The shared mailbox exists, the forwarding rule works, and the IMAP poller still creates tickets from forwarded requester email.",
        "The Entra application has SMTP.SendAsApp application permission with administrator consent.",
        "The Exchange service principal exists and the shared mailbox shows SendAs for that service principal.",
        "Authenticated SMTP policy permits the shared mailbox and the tenant authentication policy does not block OAuth SMTP submission.",
        "A direct SMTP XOAUTH2 probe authenticates successfully to smtp.office365.com:587 and sends a test message from the shared mailbox to the external test recipient.",
        "The test message's visible From and Reply-To are the O365 shared mailbox, not Gmail.",
        "Creating a ticket sends one acknowledgement to the requester, and the email contains the expected ticket number.",
        "Moving a ticket to RESOLVED sends one requester email containing the resolution note.",
        "Moving a RESOLVED ticket to CLOSED sends one requester status email.",
        "Adding an internal note sends no requester email.",
        "Adding a requester-facing note sends one requester email and does not include internal content.",
        "A reply from the external test recipient to the shared mailbox is forwarded to Gmail, is read by IMAP, and is attached to the correct ticket.",
        "An outbound notification is not processed as a new request by the IMAP path. If the forwarding configuration copies sent messages into Gmail, add the O365 shared mailbox to the inbound self-addressed sender filter before enabling the live loop.",
        "The application remains healthy after a failed SMTP connection; the ticket transition still succeeds and the mailer records a sanitized delivery failure.",
        "A secret rotation or application restart obtains a new access token without requiring an interactive administrator login.",
    ])

    add_heading(document, "7.1 Test matrix", 2)
    add_table(document, ["Test", "Expected result", "Evidence to retain"], [
        ["SMTP transport disabled", "Console transport remains available; no live SMTP call is attempted.", "Server log shows console notification."],
        ["SMTP enabled but secret absent", "Transport reports configuration failure; ticket operation is not blocked.", "Sanitized configuration error without secret."],
        ["SMTP authentication succeeds", "Ticket notification is accepted by the SMTP server.", "Live probe result and recipient inbox."],
        ["SMTP authentication fails", "Ticket update still succeeds; error is logged as a mail delivery failure.", "Log and ticket state."],
        ["RESOLVED transition", "Requester receives resolved subject and resolution note.", "Inbound test inbox."],
        ["CLOSED transition", "Requester receives closed status update.", "Inbound test inbox."],
        ["Requester reply", "Forwarded message becomes a comment or reopens a closed/resolved ticket as designed.", "Ticket timeline and IMAP result."],
        ["Application restart", "New access token is obtained; notifications continue.", "Startup and live send logs."],
    ], [1.8, 3.2, 2.0])

    add_heading(document, "8. Operational runbook", 1)
    add_heading(document, "8.1 First-time deployment", 2)
    add_number(document, "Complete the Microsoft 365 configuration in section 5.")
    add_number(document, "Implement and test the SMTP modules described in section 6.")
    add_number(document, "Set the SMTP secrets in the target deployment.")
    add_number(document, "Restart the API and confirm the startup log says SMTP is enabled without printing any secret.")
    add_number(document, "Run the live SMTP probe against the external test recipient.")
    add_number(document, "Run the ticket lifecycle test: acknowledge, start, resolve with a note, close, and reply.")
    add_number(document, "Enable the production notification path only after all acceptance checks pass.")

    add_heading(document, "8.2 Routine health checks", 2)
    add_paragraph(document, "At minimum, review these items daily or weekly depending on volume:")
    add_bullet(document, "SMTP transport enabled state and last successful delivery timestamp.")
    add_bullet(document, "Authentication failures, 401 responses, 535 responses, throttling responses, and TLS failures.")
    add_bullet(document, "Tickets that changed state but have no successful corresponding notification record, if delivery tracking is implemented.")
    add_bullet(document, "Client-secret or certificate expiry dates.")
    add_bullet(document, "The shared mailbox forwarding rule and IMAP poller health.")
    add_bullet(document, "Microsoft 365 message trace entries for messages that were accepted but reported missing by recipients.")

    add_heading(document, "8.3 Secret or certificate rotation", 2)
    add_number(document, "Create the replacement credential in Entra.")
    add_number(document, "Set the replacement credential in the deployment secret store.")
    add_number(document, "Restart the deployment and run the live SMTP probe.")
    add_number(document, "After a successful send, revoke the old Entra secret or certificate.")
    add_number(document, "Record the rotation date and verify the application still functions after the old credential is revoked.")

    add_heading(document, "9. Troubleshooting", 1)
    add_table(document, ["Symptom", "Likely cause", "Action"], [
        ["535 5.7.3 Authentication unsuccessful", "Wrong tenant/client/secret, expired secret, wrong SMTP AUTH policy, or incorrect XOAUTH2 format.", "Verify the Entra application, tenant ID, client ID, secret, SMTP AUTH policy, and XOAUTH2 username. Do not switch to a shared-mailbox password."],
        ["SMTP AUTH is disabled", "Tenant or mailbox policy blocks client submission.", "Ask the Microsoft 365 administrator to review Authenticated SMTP and authentication policy. Do not broadly disable security controls without approval."],
        ["Authentication succeeds but send is rejected", "Service principal lacks SendAs, or From is not the shared mailbox.", "Check Get-RecipientPermission and set SMTP_FROM and SMTP_REPLY_TO to the exact shared-mailbox address."],
        ["The message appears to come from Gmail", "Transport or envelope sender is configured incorrectly.", "Verify SMTP_FROM, MAIL FROM, and the selected transport. Gmail must remain inbound-only."],
        ["No email arrives and no error appears", "SMTP is disabled, Graph is disabled, and the console transport is active.", "Check the startup transport selection and set SMTP_ENABLED=true only after SMTP configuration is complete."],
        ["Reply does not return to the ticket", "Reply-To is Gmail, the forwarding rule is missing, or the reply is not threaded.", "Set Reply-To to O365, verify forwarding, and check the requester's reply subject and ticket number."],
        ["Outbound email creates a new ticket", "Forwarding copied sent messages into the IMAP mailbox and the sender is not recognized as self-addressed.", "Exclude sent messages from forwarding where possible, or add the O365 shared mailbox to the IMAP self-addressed sender filter."],
        ["Token acquisition fails", "Wrong OAuth scope, missing admin consent, or Exchange service principal not registered.", "Use https://outlook.office.com/.default, confirm SMTP.SendAsApp consent, and verify New-ServicePrincipal registration."],
        ["Templating or body looks wrong", "SMTP transport is not preserving the existing mail shape or MIME encoding.", "Assert exact subject, recipients, body, From, Reply-To, and MIME headers in the SMTP unit test."],
    ], [2.1, 2.3, 2.6])

    add_heading(document, "10. Security and isolation requirements", 1)
    add_paragraph(document, "For each isolated entity deployment:")
    add_bullet(document, "Use a separate Entra application registration and credential where practical, or document the approved shared application model and restrict its service-principal SendAs permissions to that entity's mailbox.")
    add_bullet(document, "Use a separate database, JWT secret, SMTP client secret, domain, and inbound mailbox configuration.")
    add_bullet(document, "Grant the application SendAs only to the shared mailbox for that entity.")
    add_bullet(document, "Never expose SMTP secrets, access tokens, or OAuth client secrets through API responses, frontend code, health endpoints, audit events, or log output.")
    add_bullet(document, "Restrict outbound SMTP to smtp.office365.com:587 and validate the server certificate. Do not disable TLS certificate validation in production.")
    add_bullet(document, "Review who has Microsoft 365 administrator, Exchange Online, Entra application, deployment secret, and database privileges.")
    add_bullet(document, "Keep the existing IMAP and Graph settings independent. Graph can remain disabled while SMTP is enabled.")

    add_heading(document, "11. Implementation acceptance criteria", 1)
    add_paragraph(document, "The feature is complete only when all of the following are true:")
    add_checklist(document, [
        "SMTP is selectable through environment configuration and takes precedence over console fallback when enabled.",
        "The server uses STARTTLS on smtp.office365.com:587 and SASL XOAUTH2.",
        "The application obtains an Exchange SMTP token using a dedicated Entra application and does not call Microsoft Graph for sending.",
        "The Exchange service principal has SendAs on the target O365 shared mailbox.",
        "The message is sent with the O365 shared mailbox as From, envelope sender, and Reply-To.",
        "Requester acknowledgement, status, resolution, close, and public reply emails all use the SMTP transport.",
        "Internal notes never reach requester messages.",
        "A failed mail send cannot roll back a ticket lifecycle change and is recorded as a sanitized delivery failure.",
        "The forwarding-to-Gmail IMAP path continues to ingest requester replies correctly.",
        "Tests prove the transport without logging secrets or sending live mail unless explicitly opted in.",
        "A real external test recipient receives a test message and a full ticket lifecycle notification set.",
        "The configuration and runbook contain no real secrets or credentials.",
    ])

    add_heading(document, "12. Reference commands and source map", 1)
    add_heading(document, "12.1 Microsoft 365 reference operations", 2)
    add_code(document, "# Entra application: record Tenant ID, Application (client) ID, and credential\n# API permissions: Office 365 Exchange Online / Application / SMTP.SendAsApp / Admin consent\n\nConnect-ExchangeOnline -Interactive\nGet-TransportConfig | Format-List SmtpClientAuthenticationDisabled\nGet-CASMailbox -Identity $SharedMailbox | Format-List PrimarySmtpAddress, SmtpClientAuthenticationDisabled\nGet-ServicePrincipal -Identity $ClientId\nGet-RecipientPermission -Identity $SharedMailbox | Format-Table -AutoSize")
    add_heading(document, "12.2 TicketDesk source locations", 2)
    add_table(document, ["Path", "Role in this design"], [
        ["server/src/mailer.js:67-72", "Current Graph-or-console transport selection."],
        ["server/src/mailer.js:82-167", "Safe-send wrapper and notification entry points that must remain centralized."],
        ["server/src/email/outbound.js:184-213", "Requester status and resolution message builder."],
        ["server/routes/tickets.js:451-453", "Lifecycle transition invokes requester status notification."],
        ["server/src/imap/mailService.js:205-274", "Existing forwarded-mail IMAP polling path."],
        ["server/src/imap/config.js:14-73", "Existing Gmail/IMAP environment configuration."],
        ["server/package.json:65-75", "Dependencies; no SMTP client is currently present."],
    ], [2.7, 4.2])
    add_paragraph(document, "The official Microsoft protocol reference for the OAuth SMTP XOAUTH2 format and application-permission flow is the Microsoft Learn document titled Authenticate an IMAP, POP or SMTP connection using OAuth. The official shared-mailbox reference is About shared mailboxes in Microsoft 365. The official SMTP policy reference is Enable or disable SMTP AUTH in Exchange Online.")

    add_heading(document, "13. Final operational decision", 1)
    add_callout(document, "RECOMMENDED PATH:", "Keep Gmail forwarding and IMAP for inbound mail. Add a direct Microsoft 365 SMTP OAuth transport for outbound mail. This meets the requirement that notifications leave from the O365 helpdesk shared mailbox while avoiding Microsoft Graph API calls and avoiding Gmail as the sender.", OK_FILL, GREEN)
    add_paragraph(document, "The important boundary is: the existing application already knows when to send resolution and closure notifications; the missing work is the provider-specific transport and its Microsoft 365 permissions. Complete the acceptance tests with real credentials before enabling the transport in production.")

    document.save(OUTPUT)

    check = Document(OUTPUT)
    all_text = "\n".join([paragraph.text for paragraph in check.paragraphs])
    all_text += "\n" + "\n".join(cell.text for table in check.tables for row in table.rows for cell in row.cells)
    required = [
        "smtp.office365.com",
        "SMTP.SendAsApp",
        "XOAUTH2",
        "SMTP_SHARED_MAILBOX",
        "SMTP_REPLY_TO",
        "SMTP_PASSWORD",
        "Implementation acceptance criteria",
    ]
    missing = [item for item in required if item not in all_text]
    if missing:
        raise RuntimeError("Generated document is missing required content: " + ", ".join(missing))
    print(OUTPUT)
    print("paragraphs=%d tables=%d bytes=%d" % (len(check.paragraphs), len(check.tables), os.path.getsize(OUTPUT)))


if __name__ == "__main__":
    build_document()
