from docx import Document

d = Document(r"c:\Users\dogbeide\TICKETING SYSTEM\Helpdesk-Domain-Migration-Guide-v3.docx")
parts = [par.text for par in d.paragraphs]
for t in d.tables:
    for row in t.rows:
        for c in row.cells:
            parts.append(c.text)
text = "\n".join(parts)

print("em-dash OK:", "\u2014" in text)
print("ithelpdesk.mrsholdings.com occurrences:", text.count("ithelpdesk.mrsholdings.com"))
print("bare helpdesk.mrsholdings.com (should be 0):",
      text.count("helpdesk.mrsholdings.com") - text.count("ithelpdesk.mrsholdings.com"))
print("fly.dev refs:", text.count("ticketing-system-drifting-horizon-5610.fly.dev"))
print("PORTAL_BASE_URL mentioned:", "PORTAL_BASE_URL" in text)
print("WEBHOOK_PUBLIC_URL mentioned:", "WEBHOOK_PUBLIC_URL" in text)
